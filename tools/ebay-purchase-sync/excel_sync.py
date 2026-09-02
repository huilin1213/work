"""把抓到的 eBay 订单写进本地 Excel,增量同步、按行去重 —— 重复跑脚本/常驻监听
不会往表里写出重复行。

表格是"每个商品一行"(一个订单里有几件不同商品就有几行),订单级别的字段(卖
家、订单号、tracking number、postage/VAT/buyer protection 这几个整单只算一次
的费用)在同一订单的每一行上都重复填一份,方便你在 Excel 里直接按任意列筛
选/透视,不用再手动展开订单。
"""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ebay_scraper import EbayOrder

SHEET_NAME = "购买记录"
HEADERS = [
    "卖家名称",
    "订单日期",
    "订单号",
    "物流Tracking Number",
    "产品名称",
    "产品金额",
    "产品数量",
    "VAT金额",
    "Buyer Protection Fee",
    "Postage Fee",
    "订单详情链接",
    "同步时间",
]
# 去重用的"行指纹"取这几列 —— 同一单同一件商品数量/单价都相同就认为是已经同步
# 过的行,不会重复插入(但换了单价/数量会被当成"新的一行"补进去,不会覆盖旧行,
# 保留历史痕迹,人工核对更放心)。
# 注意:金额/数量这两列必须走 _normalize_key_value 统一格式化再比较 —— xlsx 不
# 区分"整数"和"小数点后是 0 的浮点数"(比如 25.0 存进去、读回来会变成 int 25),
# 直接 str() 两边格式不一致会导致同一行被误判成"新行"重复写入。
DEDUP_COLUMNS = ["订单号", "产品名称", "产品数量", "产品金额"]
_MONEY_COLUMNS = {"产品金额"}
_INT_COLUMNS = {"产品数量"}


def _normalize_key_value(column: str, value) -> str:
    if value is None:
        return ""
    if column in _MONEY_COLUMNS:
        return f"{float(value):.2f}"
    if column in _INT_COLUMNS:
        return str(int(float(value)))
    return str(value).strip()


def _ensure_workbook(path: Path) -> tuple[Workbook, Worksheet]:
    if path.exists():
        wb = load_workbook(path)
        if SHEET_NAME in wb.sheetnames:
            ws = wb[SHEET_NAME]
        else:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append(HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_NAME
        ws.append(HEADERS)

    for col_idx, width in enumerate([18, 14, 20, 22, 32, 12, 10, 12, 18, 12, 40, 18], start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = width
    ws.freeze_panes = "A2"
    return wb, ws


def _existing_keys(ws: Worksheet) -> set[tuple]:
    header_row = [c.value for c in ws[1]]
    idx = {name: header_row.index(name) for name in DEDUP_COLUMNS if name in header_row}
    keys: set[tuple] = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        keys.add(
            tuple(
                _normalize_key_value(name, row[idx[name]] if idx[name] < len(row) else None)
                for name in DEDUP_COLUMNS
            )
        )
    return keys


def _fmt_money(v: Decimal) -> float:
    return float(v)


def sync(orders: list[EbayOrder], excel_path: str | Path) -> tuple[int, int]:
    """返回 (新增行数, 跳过的订单数)。"""
    path = Path(excel_path)
    wb, ws = _ensure_workbook(path)
    existing = _existing_keys(ws)

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    added = 0
    skipped_orders = 0

    for order in orders:
        order_added = 0
        for li in order.line_items:
            key = tuple(
                _normalize_key_value(name, raw)
                for name, raw in zip(
                    DEDUP_COLUMNS, [order.order_number, li.product_name, li.quantity, li.item_price]
                )
            )
            if key in existing:
                continue
            ws.append(
                [
                    order.seller_name,
                    order.order_date,
                    order.order_number,
                    order.tracking_number,
                    li.product_name,
                    _fmt_money(li.item_price),
                    li.quantity,
                    _fmt_money(order.vat_amount),
                    _fmt_money(order.buyer_protection_fee),
                    _fmt_money(order.postage_fee),
                    order.detail_url,
                    now,
                ]
            )
            existing.add(key)
            added += 1
            order_added += 1
        if order_added == 0:
            skipped_orders += 1

    if added:
        wb.save(path)
    return added, skipped_orders
